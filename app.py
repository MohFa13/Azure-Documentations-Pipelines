import streamlit as st
import zipfile, os, json, re, shutil
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_table_borders(table):
    tbl = table._element
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ["top","left","bottom","right","insideH","insideV"]:
        border = OxmlElement(f"w:{border_name}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")
        tblBorders.append(border)
    tbl.tblPr.append(tblBorders)

def strip_numbers(name):
    if not name: return name
    return re.sub(r'(\D+)[0-9]+$', r'\1', name)

def beautify_command(cmd: str) -> str:
    c = cmd.strip()
    if "derive" in c.lower():
        return f"Derived Column: {c}"
    if "regex" in c.lower():
        return f"Regex Command: {c}"
    if "if" in c.lower():
        return f"Conditional Logic: {c}"
    if "join" in c.lower():
        return f"Join: {c}"
    return c

def parse_zip(zip_bytes, screenshots=None):
    temp_zip = "temp_uploaded.zip"
    with open(temp_zip, "wb") as f:
        f.write(zip_bytes)
    extract_dir = "extracted"
    if os.path.exists(extract_dir):
        shutil.rmtree(extract_dir)
    os.makedirs(extract_dir, exist_ok=True)
    with zipfile.ZipFile(temp_zip, 'r') as z:
        z.extractall(extract_dir)

    pipeline_jsons, dataflow_jsons = [], []
    for root, _, files in os.walk(extract_dir):
        for fn in files:
            if fn.lower().endswith(".json"):
                fullp = os.path.join(root, fn)
                normalized = fullp.replace("\\", "/")
                if "/pipeline/" in normalized:
                    pipeline_jsons.append(fullp)
                if "/dataflow/" in normalized:
                    dataflow_jsons.append(fullp)

    dataflows_info = {}
    for df in dataflow_jsons:
        with open(df, 'r', encoding='utf-8') as f:
            dj = json.load(f)
        df_name = dj.get("name")
        tp = dj.get("properties", {}).get("typeProperties", {})
        script_lines = tp.get("scriptLines", [])
        sinks = []
        if "sinks" in tp:
            for s in tp["sinks"]:
                sinks.append(s.get("name"))
        transformations = [beautify_command(line) for line in script_lines if isinstance(line,str)]
        excel_paths = []
        for line in script_lines:
            if isinstance(line,str) and "wildcardPaths" in line:
                match = re.findall(r"'([^']+\.xlsx)'", line)
                excel_paths.extend(match)
        dataflows_info[df_name] = {
            "sinks": sinks,
            "transformations": transformations,
            "excel_paths": excel_paths
        }

    doc = Document()
    doc.add_heading("Pipelines Documentation", 0).alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    pipeline_names = []

    for pf in pipeline_jsons:
        with open(pf, 'r', encoding='utf-8') as f:
            pj = json.load(f)
        pipeline_name = pj.get("name")
        pipeline_names.append(pipeline_name)

        acts = pj.get("properties", {}).get("activities", [])
        doc.add_heading(f"Pipeline: {pipeline_name}", level=1)
        doc.add_paragraph(f"Last Publish Time: {pj.get('properties',{}).get('lastPublishTime','')}")

        doc.add_heading("Sources and Sinks", level=2)
        table = doc.add_table(rows=1, cols=5)
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text, hdr[4].text = (
            "Activity Name","Dataset Name","Type","Location","Format"
        )
        for act in acts:
            tp = act.get("typeProperties",{})
            ds_name = ""
            if act.get("inputs"):
                ds_name = strip_numbers(act["inputs"][0].get("referenceName"))
            typ = tp.get("source",{}).get("type","")
            location = ""
            if act.get("outputs"):
                outp = act["outputs"][0]
                if "parameters" in outp and "filename" in outp["parameters"]:
                    location = outp["parameters"]["filename"]
            fmt = tp.get("sink",{}).get("type","")
            row = table.add_row().cells
            row[0].text = act.get("name","")
            row[1].text = ds_name
            row[2].text = typ
            row[3].text = location
            row[4].text = fmt

        for df in dataflows_info.values():
            for ep in df["excel_paths"]:
                row = table.add_row().cells
                row[0].text = "Excel Source"
                row[1].text = os.path.basename(ep).replace(".xlsx","")
                row[2].text = "Excel"
                row[3].text = ep
                row[4].text = "xlsx"

        set_table_borders(table)

        doc.add_heading("Pipeline Queries", level=2)
        for act in acts:
            q = act.get("typeProperties",{}).get("source",{})
            query = q.get("oracleReaderQuery") or q.get("sqlReaderQuery") or q.get("query")
            if query:
                doc.add_paragraph(f"Activity: {act.get('name')}")
                pre = doc.add_paragraph()
                run = pre.add_run(query)
                run.font.name = "Consolas"
                run.font.size = Pt(9)

        doc.add_heading("Dataflow Information", level=2)
        for df_name, df in dataflows_info.items():
            doc.add_paragraph(f"Dataflow: {df_name}")
            if df["sinks"]:
                doc.add_paragraph("Sinks: " + ", ".join(df["sinks"]))
            if df["transformations"]:
                doc.add_paragraph("Transformations / Commands:")
                for t in df["transformations"]:
                    doc.add_paragraph(f"- {t}")

        if screenshots:
            doc.add_heading("Screenshots", level=2)
            for img in screenshots:
                doc.add_picture(img, width=None)

    # Decide the filename dynamically
    if pipeline_names:
        output_filename = f"{pipeline_names[0]}_documentation.docx"
    else:
        output_filename = "Pipelines_Documentation.docx"

    doc.save(output_filename)
    return output_filename

# ---------------- Streamlit App ----------------
st.title("Pipeline Documentation Generator")

uploaded = st.file_uploader("Upload your ZIP file (with pipeline/dataflow/dataset JSONs)", type=["zip"])
screenshots = None

if uploaded:
    add_screens = st.checkbox("Add screenshots at the end")
    img_files = []
    if add_screens:
        screenshots = st.file_uploader("Upload screenshots", type=["png","jpg","jpeg"], accept_multiple_files=True)
        if screenshots:
            img_files = []
            for up in screenshots:
                path = os.path.join("temp_" + up.name)
                with open(path,"wb") as f:
                    f.write(up.read())
                img_files.append(path)
            screenshots = img_files

    if st.button("Generate Documentation"):
        path = parse_zip(uploaded.read(), screenshots)
        with open(path,"rb") as f:
            st.download_button("Download DOCX", f, file_name=os.path.basename(path))
